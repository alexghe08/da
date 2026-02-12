from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

TEMPLATES_DIR = "data/templates"
if not os.path.exists(TEMPLATES_DIR):
    os.makedirs(TEMPLATES_DIR)

def create_fisa_date_template():
    doc = Document()

    # Title
    title = doc.add_heading('FIȘA DE DATE A ACHIZIȚIEI', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Section I: Contracting Authority
    doc.add_heading('I. AUTORITATEA CONTRACTANTĂ', level=1)

    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'

    cells = table.rows[0].cells
    cells[0].text = "Denumire oficială:"
    cells[1].text = "{{ autoritate_contractanta }}"

    cells = table.rows[1].cells
    cells[0].text = "Adresă:"
    cells[1].text = "{{ adresa_autoritate }}"

    cells = table.rows[2].cells
    cells[0].text = "Punct(e) de contact:"
    cells[1].text = "{{ punct_contact }}"

    cells = table.rows[3].cells
    cells[0].text = "E-mail:"
    cells[1].text = "{{ email_contact }}"

    # Section II: Object of Contract
    doc.add_heading('II. OBIECTUL CONTRACTULUI', level=1)

    doc.add_paragraph('II.1) Descriere', style='List Number')
    p = doc.add_paragraph()
    p.add_run("Titlu: ").bold = True
    p.add_run("{{ titlu_proiect }}")

    p = doc.add_paragraph()
    p.add_run("Cod CPV principal: ").bold = True
    p.add_run("{{ cod_cpv }}")

    p = doc.add_paragraph()
    p.add_run("Valoarea totală estimată: ").bold = True
    p.add_run("{{ valoare_estimata }} {{ moneda }}")

    doc.add_paragraph('II.2) Descrierea achiziției', style='List Number')
    doc.add_paragraph("{{ descriere_achizitie }}")

    # Section III: Conditions
    doc.add_heading('III. INFORMAȚII JURIDICE, ECONOMICE, FINANCIARE ȘI TEHNICE', level=1)

    doc.add_paragraph('III.1) Condiții de participare', style='List Number')
    doc.add_paragraph("Garanția de participare: {{ garantie_participare }}")

    # Save
    path = os.path.join(TEMPLATES_DIR, "fisa_date.docx")
    doc.save(path)
    print(f"Created template: {path}")

def create_contract_lucrari_template():
    doc = Document()

    # Title
    title = doc.add_heading('CONTRACT DE EXECUȚIE DE LUCRĂRI', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Nr. {{ nr_contract }} data {{ data_contract }}")

    # Preamble
    doc.add_heading('1. Părțile contractante', level=1)

    p = doc.add_paragraph()
    p.add_run("Între\n").bold = True
    p.add_run("{{ autoritate_contractanta }}")
    p.add_run(", cu sediul în {{ adresa_autoritate }}, în calitate de Achizitor, pe de o parte,\n")
    p.add_run("și\n").bold = True
    p.add_run("{{ contractant }}")
    p.add_run(", cu sediul în {{ adresa_contractant }}, în calitate de Executant, pe de altă parte,\n")
    p.add_run("au convenit încheierearea prezentului contract de execuție de lucrări.")

    # Clauses
    doc.add_heading('2. Obiectul contractului', level=1)
    p = doc.add_paragraph()
    p.add_run("2.1. Executantul se obligă să execute, să finalizeze și să întrețină lucrările pentru obiectivul: ")
    p.add_run("{{ titlu_proiect }}").bold = True
    p.add_run(", conform ofertei sale anexate.")

    doc.add_heading('3. Prețul contractului', level=1)
    p = doc.add_paragraph()
    p.add_run("3.1. Prețul convenit pentru îndeplinirea contractului, plătibil Executantului de către Achizitor, este de ")
    p.add_run("{{ valoare_contract }} {{ moneda }}").bold = True
    p.add_run(", la care se adaugă TVA.")

    doc.add_heading('4. Durata contractului', level=1)
    p = doc.add_paragraph()
    p.add_run("4.1. Durata de execuție a lucrărilor este de ")
    p.add_run("{{ durata_luni }} luni").bold = True
    p.add_run(" de la data emiterii ordinului de începere.")

    # Jinja2 Logic Example
    doc.add_heading('5. Garanția de bună execuție', level=1)
    p = doc.add_paragraph("{% if garantie_buna_executie_necesara %}")
    p.add_run("5.1. Executantul constituie garanția de bună execuție a contractului în cuantum de ")
    p.add_run("{{ procent_garantie }}%").bold = True
    p.add_run(" din prețul contractului fără TVA.")
    doc.add_paragraph("{% else %}")
    doc.add_paragraph("Nu se solicită garanție de bună execuție.")
    doc.add_paragraph("{% endif %}")

    # Save
    path = os.path.join(TEMPLATES_DIR, "contract_lucrari.docx")
    doc.save(path)
    print(f"Created template: {path}")

if __name__ == "__main__":
    create_fisa_date_template()
    create_contract_lucrari_template()
